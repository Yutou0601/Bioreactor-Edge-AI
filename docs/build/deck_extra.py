
# -*- coding: utf-8 -*-
"""補充簡報：把主簡報還沒收進去的幾件事獨立成一份。版式沿用同一個模板。

收錄四頁：
  1 「拿 k 換 r_b」的機制——主簡報只給結論，沒給機制
  2  為什麼有兩成的段算出負值——主簡報把這個弱點畫出來卻沒有回答
  3  這個速率可以拿來做什麼——全篇證明「做得到」，沒說「要幹嘛」
  4  貢獻是什麼——學術場合必問，主簡報沒有一頁講

⚠ 第 2 頁的結論不如預期：模擬在真值 0.0126 下只產生 4% 負值，
  實測卻是 21%。差距沒有被解釋，頁面照實寫，不硬掰。
"""
import os
import sys

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

import build_deck as B
from build_deck import (BLUE, BOXBG, GREY, INK, RED, RULE, THEAD, TROW2,
                        W, bullets, caption, chrome, note, page, pic, put,
                        rect, table, textbox)
import fig_concept as F

OUT = os.path.join(B.HERE, 'decks', '補充簡報_2026-08-16.pptx')
plt.rcParams['font.sans-serif'] = ['Microsoft JhengHei', 'DFKai-SB']
plt.rcParams['axes.unicode_minus'] = False
plt.rcParams['font.size'] = 11


# ── 圖：把 k 釘在不同值，看線性項被迫扛多少 ────────────────
def fig_trade():
    # ⚠ build_deck 的顏色是 PPTX 的 RGBColor，matplotlib 不吃；用色碼字串
    BLUE, RED, GREY = F.BLUE, F.RED, F.GREY
    T, k0, amp, rb0 = 13.0, 0.10, 0.31, 0.0126
    t = np.arange(0, T, F.DT)
    a0 = (amp - rb0 * T) / (1 - np.exp(-k0 * T))
    y = F.model(t, 0.63, a0, k0, rb0)
    ks = np.array([0.02, 0.035, 0.05, 0.07, 0.10, 0.15, 0.30, 0.50, 0.80])
    rb, expo, lin = [], [], []
    for kf in ks:
        e = np.exp(-kf * t)
        X = np.column_stack([np.ones_like(t), e, -t])
        c, *_ = np.linalg.lstsq(X, y, rcond=None)
        rb.append(c[2]); expo.append(c[1] * (1 - np.exp(-kf * T)))
        lin.append(c[2] * T)
    rb, expo, lin = map(np.array, (rb, expo, lin))

    fig, ax = plt.subplots(1, 2, figsize=(11.6, 4.2))
    ax[0].axhline(0, color=GREY, lw=1.0, ls=':')
    ax[0].plot(ks, rb, 'o-', lw=2.6, ms=8, color=BLUE)
    ax[0].plot([k0], [rb0], '*', ms=20, color=RED, zorder=5)
    ax[0].annotate('真相\nk=0.10, r$_b$=0.0126', xy=(k0, rb0),
                   xytext=(0.16, -0.012), fontsize=10.5, color=RED,
                   arrowprops=dict(arrowstyle='->', color=RED, lw=1.4))
    ax[0].set_xscale('log')
    ax[0].set_xticks([0.02, 0.05, 0.1, 0.3, 0.8])
    ax[0].set_xticklabels(['0.02', '0.05', '0.1', '0.3', '0.8'])
    # ⚠ 對數軸預設連次要刻度也標數字，會和自訂標籤疊成一團
    ax[0].tick_params(axis='x', which='minor', labelbottom=False)
    ax[0].set_xlabel('把 k 釘在這個值')
    ax[0].set_ylabel('擬合被迫給出的 r$_b$')
    ax[0].set_title('k 猜大一點，r$_b$ 就被推大', fontweight='bold')
    ax[0].grid(alpha=0.25)

    ax[1].bar(np.arange(len(ks)) - 0.19, expo, width=0.36,
              color='#7FB3D9', edgecolor='#123A57', lw=0.6,
              label='指數項扛的下降量')
    ax[1].bar(np.arange(len(ks)) + 0.19, lin, width=0.36,
              color='#E8836F', edgecolor='#7B1E10', lw=0.6,
              label='線性項扛的下降量')
    ax[1].axhline(0, color=GREY, lw=1.0)
    ax[1].set_xticks(range(len(ks)))
    ax[1].set_xticklabels(['%g' % k for k in ks], fontsize=9)
    ax[1].set_xlabel('把 k 釘在這個值')
    ax[1].set_ylabel('這一項解釋了多少下降（kg/cm²）')
    ax[1].set_title('兩項在同一次擬合裡交換工作量', fontweight='bold',
                    color=RED)
    ax[1].legend(loc='upper center', fontsize=10, framealpha=0.95)
    ax[1].grid(alpha=0.22, axis='y')
    fig.suptitle('同一段資料、同一個真相，只是把 k 釘在不同的值',
                 fontsize=13, fontweight='bold')
    fig.tight_layout(rect=(0, 0, 1, 0.92))
    q = os.path.join(F.OUT, 'figX1_trade.png')
    fig.savefig(q, dpi=180); plt.close(fig)
    print('圖：k 換 r_b → figX1_trade.png')


# ── 逐頁 ───────────────────────────────────────────────────
def cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, Cm(0.6), Cm(5.6), W - Cm(1.2), Cm(0.05), RULE)
    tf = textbox(s, Cm(2.0), Cm(6.6), W - Cm(4.0), Cm(5.0))
    put(tf, [('補　充　簡　報', 38, True, INK)], first=True,
        space_after=14, align=PP_ALIGN.CENTER)
    put(tf, [('主簡報未收錄的四件事', 20, False, INK)], space_after=12,
        align=PP_ALIGN.CENTER)
    put(tf, [('機制 ｜ 負值 ｜ 用途 ｜ 貢獻', 15, False, GREY)],
        align=PP_ALIGN.CENTER)
    rect(s, Cm(0.6), Cm(12.6), W - Cm(1.2), Cm(0.05), RULE)
    chrome(s)


def x1(prs):
    s = page(prs, '「拿 k 換 r_{b}」到底是怎麼換的',
             '主簡報只說兩者會連動，沒說為什麼')
    pic(s, 'figX1_trade.png', Cm(6.0), Cm(3.1), Cm(21.8))
    caption(s, Cm(1.6), Cm(11.4), W - Cm(3.2), '-',
            '同一段資料、同一個真相，只把 k 釘在不同的值。'
            '左：擬合被迫給出的 r_{b}。右：兩項各自解釋掉多少下降量。')
    tf = textbox(s, Cm(1.6), Cm(12.4), W - Cm(3.2), Cm(2.6))
    put(tf, [('• k 猜小 → e^{−kt} 在整個窗裡幾乎是直線 → '
              '指數項自己就吃掉整段下降 → ', 15),
             ('線性項不必做事，r_{b} 被壓小甚至變負', 15, True, BLUE)],
        first=True, space_after=7)
    put(tf, [('• k 猜大 → 指數項前面掉完就平了 → 剩下的長直段沒人扛 → ',
              15), ('線性項被迫接手，r_{b} 被推大', 15, True, RED)],
        space_after=7)
    put(tf, [('• 所以就算每一段的真實速率完全相同，'
              'k̂ 與 r̂_{b} 也會一起大一起小 —— 這就是「假影」的來源。',
              15, False, GREY)])
    note(s, '★ 實測：k 釘 0.02 給出 r_{b} = −0.031，釘 0.80 給出 +0.022，'
            '真相是 0.0126。而 k 從 0.05 移到 0.10，殘差平方和只差 '
            '1.1×10⁻⁴ —— 資料幾乎分不出誰對。')


def x2(prs):
    s = page(prs, '為什麼有兩成的段算出負的速率',
             '主簡報把這個弱點畫了出來，這頁回答它')
    table(s, Cm(1.4), Cm(3.2), W - Cm(2.8),
          [['', '負值比例', '說明'],
           ['實測（資料集 C 全部 280 段）', '21%', '物理上不可能有負的生物速率'],
           ['模擬（真值固定 0.0126）', '4%', '同樣的雜訊、量化與估計器'],
           ['差距', '17 個百分點', '⚠ 沒有被解釋']],
          widths=[Cm(12.4), Cm(6.6), Cm(12.1)], fs=14, hi_rows=(3,),
          hi_col=1)
    tf = textbox(s, Cm(1.6), Cm(8.0), W - Cm(3.2), Cm(6.0))
    put(tf, [('能講的部分', 17, True, BLUE)], first=True, space_after=9)
    bullets(tf, ['單一段的估計本來就很吵：擬合在 k 上稍微偏一點，'
                 'r_{b} 就會被推到負的（見前一頁）',
                 '所以本文從頭到尾只用中位數，不用逐段值'],
            fs=15, first=False, gap=9)
    put(tf, [('不能講的部分', 17, True, RED)], space_after=9)
    bullets(tf, ['模擬只產生 4% 負值，實測 21% —— 多出來的 17 個百分點'
                 '目前沒有解釋',
                 '可能的方向：真值本身逐段不同、部分段混入其他事件、'
                 '或模擬器仍與實測可分辨（AUC 0.695）'],
            fs=15, first=False, gap=9)
    note(s, '★ 誠實的說法：負值比例偏高，是「逐段的值不可單獨使用」'
            '這條限制的直接證據，不是方法失效的證據 —— '
            '但兩者的差距我們還沒有查清楚。')


def x3(prs):
    s = page(prs, '這個速率可以拿來做什麼',
             '主簡報證明了「做得到」，這頁講「要幹嘛」')
    items = [('看菌還活著嗎', '速率掉下來就是活性衰退的訊號，'
              '不必等到甲烷濃度掉才發現'),
             ('比較操作條件', '換液、換循環時間、換氣體比例前後，'
              '有一把共同的尺可以比'),
             ('不打擾生產', '不加感測器、不對槽體施加激發，'
              '用的是設備本來就會產生的紀錄'),
             ('跑得動', '純 NumPy，塞得進監控電腦的 60 MB 預算，'
              '不必上傳雲端')]
    cw, gap = Cm(7.5), Cm(0.65)
    for i, (t1, t2) in enumerate(items):
        x = Cm(1.5) + i * (cw + gap)
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Cm(3.4),
                               cw, Cm(5.2))
        b.fill.solid(); b.fill.fore_color.rgb = BOXBG
        b.line.color.rgb = THEAD; b.line.width = Pt(1.0)
        b.shadow.inherit = False
        tb = textbox(s, x + Cm(0.35), Cm(3.85), cw - Cm(0.7), Cm(4.4))
        put(tb, [(t1, 16, True, BLUE)], first=True, space_after=9,
            align=PP_ALIGN.CENTER)
        put(tb, [(t2, 12.5, False, GREY)], align=PP_ALIGN.CENTER)
    tf = textbox(s, Cm(1.6), Cm(9.4), W - Cm(3.2), Cm(4.6))
    put(tf, [('前提是把它當「相對指標」用。', 17, True, RED)],
        first=True, space_after=11)
    put(tf, [('本文報的是壓力單位的速率，還沒換算成 mol/hr'
              '（缺頭空體積與操作溫度）；也還沒有文獻或現場基準'
              '可以說 0.0126 算快還是慢。', 15, False, GREY)],
        space_after=11)
    put(tf, [('所以現階段能做的是「', 15),
             ('同一台機器、不同時間互相比', 15, True, BLUE),
             ('」，不是拿去跟別人的反應器比。', 15)])
    note(s, '★ 待補：把速率換成 mol/hr 需要頭空體積與操作溫度（問洪博）；'
            '要判斷快慢需要文獻基準（需 DOI 查證）。'
            '這兩件事補上之前，不要在報告裡宣稱絕對水準。')


def x4(prs):
    s = page(prs, '貢獻是什麼 —— 相對於前人',
             '學術場合必問，主簡報沒有一頁講')
    table(s, Cm(1.4), Cm(3.2), W - Cm(2.8),
          [['項目', '前人已有', '本文'],
           ['統計方法', '間接推論、經驗校準、分類器雙樣本檢定'
            '皆為既有技術', '沿用，未宣稱方法新穎'],
           ['監控控制器層', 'Iturbe 等已做「監控控制器以區分異動來源」',
            '讀不到控制器，只能由壓力反推'],
           ['動態除氣法', 'Bandyopadhyay 1967 即有，早 59 年',
            '不宣稱是新的量測原理'],
           ['所報的量', '耦合式寫法下，r_{b} 只以 P_{eq} − r_{b}/k '
            '出現，不可辨識', '★ 改以相加式分解，明確定義為線性項的係數'],
           ['真正新的地方', '厭氧消化領域尚無「單一量化壓力通道 +'
            '設計配對校準」的作法', '★ 領域缺口']],
          widths=[Cm(6.4), Cm(14.0), Cm(10.7)], fs=13, hi_rows=(4,),
          hi_col=2)
    tf = textbox(s, Cm(1.6), Cm(9.6), W - Cm(3.2), Cm(4.4))
    put(tf, [('把話講白：', 17, True, INK),
             ('本文沒有發明新的統計方法。', 17, True, RED)],
        first=True, space_after=11)
    put(tf, [('貢獻在於把既有的校準邏輯，用在一個'
              '「不能加裝、不能激發、只有一支量化壓力計」的生產設備上，'
              '並且把做不到的部分一起報出來。', 15)], space_after=11)
    put(tf, [('⚠ 不要說「我們誠實報告了失敗的回收」：', 15, True, RED),
             ('8 月版的 p = 0.40 求的是耦合式裡的生物消耗率，'
              '本文求的是相加式裡線性項的係數 —— ', 15),
             ('兩者本來就是不同的量，不是同一件事沒做成。', 15, True,
              RED)])
    note(s, '⚠ 這頁的每一條都應該在投稿前再查一次文獻，'
            '尤其「領域尚無此作法」這種宣稱最容易被審稿人推翻。')


def main():
    fig_trade()
    prs = Presentation()
    prs.slide_width, prs.slide_height = B.W, B.H
    cover(prs)
    for fn in (x1, x2, x3, x4):
        fn(prs)
    prs.save(OUT)
    print('OK', OUT, '共', len(prs.slides._sldIdLst), '頁')


if __name__ == '__main__':
    main()
