
# -*- coding: utf-8 -*-
"""簡報版面稽核：重疊、超界、文字溢出。

之前只驗「有沒有超出版面」，結果漏掉了真正會被看見的問題——
圖片壓在圖說上、文字框裝不下自己的內容。這支把三件事都驗：

  1. 超界      任何元件跑出版面或壓進底部色帶
  2. 重疊      兩個元件的矩形相交（刻意的「文字框放在色塊裡」除外）
  3. 文字溢出  依字級與字數估算實際需要的高度，超過宣告高度就報

第 3 項是估算：中文字寬約 1 em、西文約 0.5 em，行高 1.32 倍。
估得保守，寧可多報也不要漏報。
"""
import sys

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Cm

EMU = 360000.0                       # 1 cm
import os

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    'decks', '報告簡報_2026-08-14_定版論文.pptx')
LINE = 1.32                          # 行高倍率
PT2CM = 0.03528


def cm(v):
    return (v or 0) / EMU


def text_cm(tf, width_cm):
    """估算文字實際需要的高度（cm）。"""
    total, paras = 0.0, list(tf.paragraphs)
    for p in paras:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        sz = max((r.font.size.pt if r.font.size else 12) for r in runs)
        em = sz * PT2CM
        w = sum((em if ord(c) > 0x2E80 else em * 0.52)
                for r in runs for c in r.text)
        lines = max(1, -(-w // max(width_cm, 0.1)))   # 無條件進位
        total += lines * em * LINE
        if p is not paras[-1]:          # 最後一段的段後距不佔可見高度
            total += (p.space_after.pt if p.space_after else 0) * PT2CM
    return total


def box(sh):
    """回傳 (左, 上, 右, 下)，單位 cm。表格用實際列高欄寬。"""
    if sh.has_table:
        w = sum(c.width for c in sh.table.columns)
        h = sum(r.height for r in sh.table.rows)
    else:
        w, h = sh.width, sh.height
        if sh.has_text_frame and not sh.shape_type == 13:
            h = max(h, text_cm(sh.text_frame, cm(w)) * EMU)
    return cm(sh.left), cm(sh.top), cm(sh.left) + cm(w), cm(sh.top) + cm(h)


def inside(a, b, pad=0.12):
    """a 是否幾乎完全落在 b 之內（刻意把文字放進色塊的情形）。"""
    return (a[0] >= b[0] - pad and a[1] >= b[1] - pad
            and a[2] <= b[2] + pad and a[3] <= b[3] + pad)


def main():
    prs = Presentation(PATH)
    W, H = cm(prs.slide_width), cm(prs.slide_height)
    BAR = H - 0.85
    bad = 0

    for i, s in enumerate(prs.slides, 1):
        items = []
        for sh in s.shapes:
            if sh.left is None or sh.top is None:
                continue
            b = box(sh)
            full = (b[2] - b[0]) >= W - 0.05          # 色帶、細線
            items.append((sh, b, full))

        for sh, b, full in items:
            if full or b[1] >= BAR:
                continue
            if b[2] > W + 0.03 or b[0] < -0.03:
                print('  ✘ p%-2d 左右超界  %s  右=%.2f'
                      % (i, sh.shape_type, b[2])); bad += 1
            if b[3] > BAR + 0.03:
                print('  ✘ p%-2d 壓到色帶  %s  底=%.2f（限 %.2f）'
                      % (i, sh.shape_type, b[3], BAR)); bad += 1

        for m in range(len(items)):
            for n in range(m + 1, len(items)):
                (s1, b1, f1), (s2, b2, f2) = items[m], items[n]
                if f1 or f2 or b1[1] >= BAR or b2[1] >= BAR:
                    continue
                ox = min(b1[2], b2[2]) - max(b1[0], b2[0])
                oy = min(b1[3], b2[3]) - max(b1[1], b2[1])
                if ox <= 0.06 or oy <= 0.06:
                    continue
                if inside(b1, b2) or inside(b2, b1):
                    continue                          # 文字放在色塊裡
                t1 = (s1.text_frame.text[:16] if s1.has_text_frame else
                      ('圖' if s1.shape_type == 13 else '形'))
                t2 = (s2.text_frame.text[:16] if s2.has_text_frame else
                      ('圖' if s2.shape_type == 13 else '形'))
                print('  ✘ p%-2d 重疊 %.2f×%.2f cm　[%s] × [%s]'
                      % (i, ox, oy, t1, t2)); bad += 1

    print('\n投影片 %d　問題 %d 處' % (len(prs.slides._sldIdLst), bad))
    return bad


if __name__ == '__main__':
    sys.exit(1 if main() else 0)
