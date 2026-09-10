
# -*- coding: utf-8 -*-
"""
把**真向量** SVG 內嵌進 .docx（Word 2016+ 原生支援）
════════════════════════════════════════════════════════════════════════

**為什麼要自己寫**：Word 的向量格式是 EMF，而本機沒有 inkscape /
libreoffice / imagemagick，matplotlib 也寫不出 EMF。可行解是 SVG——
Word 2016 起原生支援，但 `python-docx` 不認得 SVG（它的 image 模組會
嗅探檔頭，只收 PNG/JPEG/GIF/BMP/TIFF）。

**OOXML 的做法**（Microsoft 2016 SVG 擴充）：
  1. 照常插入一張 **PNG 後備圖**——舊版 Word 與非 Word 檢視器看這張
  2. 另外把 SVG 當成一個 image part 加進 package
  3. 在 `<a:blip>` 底下塞一段 `<a:extLst>`，用 `asvg:svgBlip` 指向 SVG

  支援 SVG 的 Word 會改用向量版算繪；不支援的自動退回 PNG。
  兩者都在檔案裡，投稿端不論用什麼開都不會壞。

用法：
    from svg_docx import add_vector_picture
    add_vector_picture(doc, 'fig1.svg', 'fig1.png', width=Inches(4.8))
"""
import os

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn, nsmap
from docx.shared import Inches                                     # noqa: F401

SVG_NS = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
SVG_EXT_URI = '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'
# 這個 URI 是微軟為 svgBlip 定義的固定值，不可改。


def _next_svg_partname(package):
    n = 1
    existing = {p.partname for p in package.iter_parts()}
    while PackURI(f'/word/media/vector{n}.svg') in existing:
        n += 1
    return PackURI(f'/word/media/vector{n}.svg')


def add_vector_picture(doc_or_run, svg_path, png_path, width=None):
    """插入一張圖：PNG 為後備、SVG 為向量主體。回傳該 run。

    `doc_or_run` 可以是 Document（會自己開新段落）或既有的 run。
    """
    # ⚠ 照片（equipment.png）沒有、也不該有 SVG 版：向量化一張相片沒有
    #   意義。SVG 缺席時只放 PNG，不再視為錯誤；PNG 缺席才是真的壞掉。
    if not os.path.exists(png_path):
        raise FileNotFoundError(png_path)
    has_svg = os.path.exists(svg_path)

    if hasattr(doc_or_run, 'add_paragraph'):
        run = doc_or_run.add_paragraph().add_run()
    else:
        run = doc_or_run

    # ① PNG 後備——交給 python-docx 正常處理，尺寸也由它算
    run.add_picture(png_path, width=width)
    if not has_svg:
        return run                     # 相片：到此為止，沒有向量層可掛

    # ② 把 SVG 當成 image part 塞進 package
    doc_part = run.part
    with open(svg_path, 'rb') as fh:
        svg_bytes = fh.read()
    partname = _next_svg_partname(doc_part.package)
    svg_part = Part(partname, 'image/svg+xml', svg_bytes, doc_part.package)
    rid = doc_part.relate_to(svg_part, RT.IMAGE)

    # ③ 在剛插入的 blip 底下掛 svgBlip
    blips = run._r.findall('.//' + qn('a:blip'))
    if not blips:
        raise RuntimeError('找不到 a:blip，PNG 可能沒插進去')
    blip = blips[-1]

    from docx.oxml import parse_xml
    ext_xml = (
        f'<a:extLst xmlns:a="{nsmap["a"]}" '
        f'xmlns:r="{nsmap["r"]}" xmlns:asvg="{SVG_NS}">'
        f'<a:ext uri="{SVG_EXT_URI}">'
        f'<asvg:svgBlip r:embed="{rid}"/>'
        f'</a:ext></a:extLst>'
    )
    blip.append(parse_xml(ext_xml))
    return run


def selftest(outdir):
    """產生一個最小測試檔並驗證封裝結果。"""
    import zipfile
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from docx import Document

    fig, ax = plt.subplots(figsize=(3, 1.6))
    ax.plot([0, 1, 2, 3], [0, 1, 0.5, 1.6], lw=1.2)
    ax.set_xlabel('t'); ax.set_ylabel('P')
    svg = os.path.join(outdir, '_svgtest.svg')
    png = os.path.join(outdir, '_svgtest.png')
    fig.savefig(svg); fig.savefig(png, dpi=600)
    plt.close(fig)

    doc = Document()
    doc.add_paragraph('vector embed test')
    add_vector_picture(doc, svg, png, width=Inches(3.0))
    out = os.path.join(outdir, '_svgtest.docx')
    doc.save(out)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        has_svg = [n for n in names if n.endswith('.svg')]
        doc_xml = z.read('word/document.xml').decode('utf-8')
        ct = z.read('[Content_Types].xml').decode('utf-8')
    ok_part = bool(has_svg)
    ok_ref = 'svgBlip' in doc_xml
    ok_ct = 'image/svg+xml' in ct
    print(f'   SVG part 存在      {has_svg}   {"✓" if ok_part else "✘"}')
    print(f'   document.xml 有 svgBlip   {"✓" if ok_ref else "✘"}')
    print(f'   Content_Types 有 svg+xml  {"✓" if ok_ct else "✘"}')
    for f in (svg, png, out):
        os.remove(f)
    return ok_part and ok_ref and ok_ct


if __name__ == '__main__':
    import sys
    d = os.path.dirname(os.path.abspath(__file__))
    print('══ SVG→docx 內嵌自我測試 ══')
    sys.exit(0 if selftest(d) else 1)
